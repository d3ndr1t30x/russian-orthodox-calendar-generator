from datetime import date, timedelta

from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader

from orthodox_calendar.models import CalendarDay, FastLevel, Fasting, Saint, ServiceRank, ServiceRankInfo
from orthodox_calendar.rendering.docx_renderer import DocxRenderer
from orthodox_calendar.rendering.pdf_renderer import IconRenderer, PdfOptions, PdfRenderer
from orthodox_calendar.service_ranks import rank_text_is_red, symbol_colour_for, symbol_for


def sample_day() -> CalendarDay:
    civil = date(2027, 1, 3)
    day = CalendarDay(
        civil,
        civil - timedelta(days=13),
        saints=[Saint(42, "Reference", "Reference-ranked saint", civil, selected=True, service_rank=ServiceRank.POLYELEOS)],
        fasting=Fasting(FastLevel.FISH, "Fast", "Fish permitted"),
        liturgical_week="32nd Week after Pentecost",
        tone=6,
        service_rank=ServiceRankInfo(ServiceRank.POLYELEOS),
        default_primary_saint_id="id:42",
        primary_saint_id="id:42",
    )
    return day


def test_exact_reference_rank_symbols_colours_and_no_sign():
    assert symbol_for(ServiceRank.GREAT_FEAST) == "🕀"
    assert symbol_for(ServiceRank.VIGIL) == "🕁"
    assert symbol_for(ServiceRank.POLYELEOS) == "🕂"
    assert symbol_for(ServiceRank.DOXOLOGY) == symbol_for(ServiceRank.SIX_STICHERA) == "🕃"
    assert symbol_for(ServiceRank.NO_SIGN) == ""
    assert symbol_colour_for(ServiceRank.DOXOLOGY) == "C00000"
    assert symbol_colour_for(ServiceRank.SIX_STICHERA) == "000000"
    assert rank_text_is_red(ServiceRank.VIGIL) and rank_text_is_red(ServiceRank.POLYELEOS)
    assert not rank_text_is_red(ServiceRank.DOXOLOGY)


def test_saints_and_source_fasting_symbol_choice_default_to_reference_behaviour():
    assert not Saint(1, "Saint", "Saint", date(2027, 1, 1)).selected
    assert IconRenderer.fasting_symbol_names(Fasting(FastLevel.FISH, "Fast", "Fish permitted")) == ["fish"]
    assert IconRenderer.fasting_symbol_names(Fasting(FastLevel.WINE_OIL, "Fast", "Wine and oil permitted")) == ["oil"]
    assert IconRenderer.fasting_symbol_names(Fasting(FastLevel.STRICT, "Strict fast", "No food")) == []


def test_docx_uses_inline_rank_top_right_fast_symbol_eight_point_type_and_grid_legend(tmp_path):
    output = tmp_path / "reference-symbols.docx"
    DocxRenderer().render(output, [sample_day()], PdfOptions(2027, "Queensland", months=[1]))
    document = Document(output)
    assert document.styles["Normal"].font.size.pt == 8
    cells = [cell for row in document.tables[0].rows for cell in row.cells]
    day_cell = next(cell for cell in cells if "Reference-ranked saint" in cell.text)
    paragraphs = [paragraph for paragraph in day_cell.paragraphs if paragraph.text]
    date_paragraph = next(paragraph for paragraph in paragraphs if "🐟" in paragraph.text)
    week_index = next(index for index, paragraph in enumerate(paragraphs) if "32nd Week after Pentecost" in paragraph.text)
    saint_index = next(index for index, paragraph in enumerate(paragraphs) if "Reference-ranked saint" in paragraph.text)
    saint_paragraph = paragraphs[saint_index]
    assert week_index < saint_index
    assert saint_paragraph.text.startswith("🕂 Reference-ranked saint")
    assert "Polyeleos" not in day_cell.text and "Fish permitted" not in day_cell.text
    symbol_run = next(run for run in saint_paragraph.runs if run.text == "🕂")
    assert symbol_run.font.color.rgb.__str__() == "C00000"
    assert symbol_run._element.rPr.rFonts.get(qn("w:cs")) == "Segoe UI Symbol"
    fish_run = next(run for run in date_paragraph.runs if run.text == "🐟")
    assert fish_run.font.size.pt == 14 and fish_run._element.rPr.rFonts.get(qn("w:ascii")) == "Segoe UI Emoji"
    grid_text = "\n".join(cell.text for cell in cells)
    assert "Polyeleos-ranked service" in grid_text and "Fish permitted" in grid_text
    assert all("Polyeleos-ranked service" not in paragraph.text for paragraph in document.paragraphs)


def test_week_tone_toggle_controls_both_word_and_pdf(tmp_path):
    day = sample_day()
    for enabled in (True, False):
        options = PdfOptions(
            2027,
            "Queensland",
            months=[1],
            include_liturgical_week_tone=enabled,
            include_fasting_legend=False,
            include_service_rank_legend=False,
        )
        docx_path = tmp_path / f"tone-{enabled}.docx"
        pdf_path = tmp_path / f"tone-{enabled}.pdf"
        DocxRenderer().render(docx_path, [day], options)
        PdfRenderer().render(pdf_path, [day], options)
        word_text = "\n".join(cell.text for row in Document(docx_path).tables[0].rows for cell in row.cells)
        pdf_text = PdfReader(pdf_path).pages[0].extract_text() or ""
        assert ("32nd Week after Pentecost" in word_text) is enabled
        assert ("32nd Week after Pentecost" in pdf_text) is enabled
        assert "Polyeleos" not in next(cell.text for row in Document(docx_path).tables[0].rows for cell in row.cells if "Reference-ranked saint" in cell.text)
    fonts = PdfReader(tmp_path / "tone-True.pdf").pages[0]["/Resources"]["/Font"].get_object().values()
    assert any("SegoeUISymbol" in str(font.get_object().get("/BaseFont", "")) for font in fonts)
